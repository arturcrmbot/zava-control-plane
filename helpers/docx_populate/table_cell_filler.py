"""Fill empty placeholder cells in existing master tables."""
from __future__ import annotations

from docx.document import Document as DocxDocument


def fill_table_cell(doc: DocxDocument, table_index: int, row: int, col: int, value: str) -> None:
    if table_index >= len(doc.tables):
        raise IndexError(f"Table index {table_index} out of range")
    table = doc.tables[table_index]
    cell = table.cell(row, col)
    cell.text = value
