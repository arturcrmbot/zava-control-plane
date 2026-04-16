"""Resolve inline placeholders in the master docx.

Placeholders handled:
  1. `<<...>>` markers — replace entire paragraph with authored MD content
  2. `Cut and paste below from Artur's document...` — delete the placeholder paragraph
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Literal, Optional

from docx import Document
from docx.document import Document as DocxDocument

from helpers.docx_populate.md_to_docx import render_md_into_doc


@dataclass(frozen=True)
class PlaceholderRule:
    match_prefix: str
    resolution: Literal["replace_with", "delete_paragraph"]
    content: Optional[str] = None  # MD content for replace_with


def resolve_inline_placeholder(doc: DocxDocument, rule: PlaceholderRule) -> bool:
    """Apply the rule. Returns True if a matching paragraph was found and resolved."""
    needle = rule.match_prefix.strip()
    target_para = None
    for para in doc.paragraphs:
        if para.text.strip().startswith(needle):
            target_para = para
            break
    if target_para is None:
        return False

    target_el = target_para._element
    parent = target_el.getparent()

    if rule.resolution == "delete_paragraph":
        parent.remove(target_el)
        return True

    assert rule.content is not None, "replace_with requires content"
    tmp = Document()
    render_md_into_doc(tmp, rule.content)
    new_elements = [p._element for p in tmp.paragraphs]
    target_index = list(parent).index(target_el)
    for new_el in new_elements:
        parent.insert(target_index, new_el)
        target_index += 1
    parent.remove(target_el)
    return True
