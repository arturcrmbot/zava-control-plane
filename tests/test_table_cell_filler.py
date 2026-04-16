from docx import Document
from helpers.docx_populate.table_cell_filler import fill_table_cell


def test_fill_empty_cell(tmp_path, fixtures_dir):
    doc = Document(fixtures_dir / "mini-master.docx")
    table = doc.tables[0]
    assert table.cell(1, 1).text.strip() == ""
    fill_table_cell(doc, table_index=0, row=1, col=1, value="§5, §10")
    assert table.cell(1, 1).text == "§5, §10"
